import csv
import io
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..adapters import MockPlatformAdapter
from ..adapters.base import UnifiedItem, UnifiedOrder
from ..bootstrap import ensure_development_data
from ..database import get_db
from ..models import Order, OrderItem, Shop
from ..security import Principal, get_principal

router = APIRouter(prefix="/orders", tags=["统一订单中心"])


class BatchRequest(BaseModel):
    order_ids: list[str] = Field(min_length=1, max_length=1000)


def serialize_order(order: Order) -> dict:
    return {
        "id": order.id,
        "platform": order.platform,
        "shop_id": order.shop_id,
        "platform_order_no": order.platform_order_no,
        "erp_order_no": order.erp_order_no,
        "status": order.status,
        "after_sale_status": order.after_sale_status,
        "audit_status": order.audit_status,
        "print_status": order.print_status,
        "receiver_name": order.receiver_name,
        "receiver_phone_masked": order.receiver_phone_masked,
        "receiver_address_masked": order.receiver_address_masked,
        "warehouse_name": order.warehouse_name,
        "paid_amount": str(order.paid_amount),
        "discount_amount": str(order.discount_amount),
        "remark": order.remark,
        "paid_at": order.paid_at,
        "created_at": order.created_at,
        "items": [
            {
                "id": item.id,
                "platform_sku": item.platform_sku,
                "internal_sku": item.internal_sku,
                "product_name": item.product_name,
                "specification": item.specification,
                "quantity": item.quantity,
                "unit_price": str(item.unit_price),
                "barcode": item.barcode,
            }
            for item in order.items
        ],
    }


def upsert_unified(db: Session, tenant_id: str, incoming: UnifiedOrder) -> tuple[Order, bool]:
    order = db.scalar(
        select(Order).where(
            Order.tenant_id == tenant_id,
            Order.platform == incoming.platform,
            Order.shop_id == incoming.shop_id,
            Order.platform_order_no == incoming.platform_order_no,
        )
    )
    created = order is None
    if order is None:
        order = Order(
            tenant_id=tenant_id,
            shop_id=incoming.shop_id,
            platform=incoming.platform,
            platform_order_no=incoming.platform_order_no,
            erp_order_no=f"SO{incoming.created_at:%Y%m%d}{incoming.platform_order_no[-8:]}",
            status=incoming.status,
            after_sale_status=incoming.after_sale_status,
            receiver_name=incoming.receiver_name,
            receiver_phone_masked=incoming.receiver_phone_masked,
            receiver_address_masked=incoming.receiver_address_masked,
            warehouse_name="杭州一号仓",
            paid_amount=incoming.paid_amount,
            discount_amount=incoming.discount_amount,
            paid_at=incoming.paid_at,
            created_at=incoming.created_at,
            remark=incoming.remark,
        )
        for item in incoming.items:
            order.items.append(
                OrderItem(
                    tenant_id=tenant_id,
                    platform_sku=item.platform_sku,
                    internal_sku=item.internal_sku,
                    product_name=item.product_name,
                    specification=item.specification,
                    quantity=item.quantity,
                    unit_price=item.unit_price,
                    barcode=item.barcode,
                )
            )
        db.add(order)
    else:
        requires_review = (
            order.status != incoming.status
            or order.after_sale_status != incoming.after_sale_status
            or order.receiver_address_masked != incoming.receiver_address_masked
        )
        order.status = incoming.status
        order.after_sale_status = incoming.after_sale_status
        order.receiver_address_masked = incoming.receiver_address_masked
        if requires_review:
            order.audit_status = "PENDING"
    return order, created


@router.post("/mock/sync")
def sync_mock(
    count: int = Query(1000, ge=1, le=10000),
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    ensure_development_data(db)
    created_count = 0
    updated_count = 0
    for incoming in MockPlatformAdapter().fetch_orders(count):
        _, created = upsert_unified(db, principal.tenant_id, incoming)
        created_count += int(created)
        updated_count += int(not created)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "订单同步发生幂等冲突，请重试") from exc
    return {"received": count, "created": created_count, "updated": updated_count}


@router.get("")
def list_orders(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    platform: str | None = None,
    shop_id: str | None = None,
    status: str | None = None,
    keyword: str | None = None,
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    conditions = [Order.tenant_id == principal.tenant_id]
    if platform:
        conditions.append(Order.platform == platform)
    if shop_id:
        conditions.append(Order.shop_id == shop_id)
    if status:
        conditions.append(Order.status == status)
    if keyword:
        term = f"%{keyword.strip()}%"
        conditions.append(
            or_(
                Order.platform_order_no.ilike(term),
                Order.erp_order_no.ilike(term),
                Order.items.any(OrderItem.product_name.ilike(term)),
                Order.items.any(OrderItem.internal_sku.ilike(term)),
            )
        )
    total = db.scalar(select(func.count()).select_from(Order).where(*conditions)) or 0
    orders = db.scalars(
        select(Order)
        .where(*conditions)
        .order_by(Order.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    shop_names = {shop.id: shop.name for shop in db.scalars(select(Shop)).all()}
    records = []
    for order in orders:
        record = serialize_order(order)
        record["shop_name"] = shop_names.get(order.shop_id, "未知店铺")
        records.append(record)
    return {"records": records, "total": total, "page": page, "page_size": page_size}


@router.get("/{order_id}")
def order_detail(
    order_id: str,
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    order = db.scalar(
        select(Order).where(Order.id == order_id, Order.tenant_id == principal.tenant_id)
    )
    if not order:
        raise HTTPException(404, "订单不存在")
    return serialize_order(order)


@router.post("/batch/audit")
def batch_audit(
    payload: BatchRequest,
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    orders = db.scalars(
        select(Order).where(Order.tenant_id == principal.tenant_id, Order.id.in_(payload.order_ids))
    ).all()
    for order in orders:
        if order.status in {"CANCELLED", "REFUNDED"} or order.after_sale_status != "NONE":
            continue
        order.audit_status = "APPROVED"
    db.commit()
    return {"updated": sum(order.audit_status == "APPROVED" for order in orders)}


@router.post("/batch/print-queue")
def batch_print_queue(
    payload: BatchRequest,
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    orders = db.scalars(
        select(Order).where(Order.tenant_id == principal.tenant_id, Order.id.in_(payload.order_ids))
    ).all()
    for order in orders:
        order.print_status = "QUEUED"
    db.commit()
    return {"queued": len(orders)}


def _rows_from_upload(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith(".csv"):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [
            dict(zip(headers, ("" if value is None else str(value) for value in row), strict=False))
            for row in rows[1:]
        ]
    raise HTTPException(415, "仅支持 CSV 或 XLSX 文件")


@router.post("/import")
async def import_orders(
    file: UploadFile = File(...),
    principal: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
) -> dict:
    ensure_development_data(db)
    rows = _rows_from_upload(file.filename or "", await file.read())
    shop = db.scalar(select(Shop).where(Shop.tenant_id == principal.tenant_id))
    created = 0
    errors: list[dict] = []
    for line, row in enumerate(rows, start=2):
        try:
            amount = Decimal(row.get("实付金额", "0"))
            order_no = row["平台订单号"].strip()
            if not order_no:
                raise ValueError("平台订单号为空")
            incoming = UnifiedOrder(
                platform=(row.get("平台") or "CSV").upper(),
                shop_id=shop.id,
                platform_order_no=order_no,
                status=row.get("订单状态") or "PENDING_SHIPMENT",
                after_sale_status="NONE",
                receiver_name=row.get("收件人") or "未提供",
                receiver_phone_masked=row.get("手机号") or "***",
                receiver_address_masked=row.get("地址") or "未提供",
                paid_amount=amount,
                discount_amount=Decimal("0"),
                paid_at=datetime.now(UTC),
                created_at=datetime.now(UTC),
                remark=row.get("备注") or "",
                items=(
                    UnifiedItem(
                        platform_sku=row.get("平台SKU") or row.get("SKU") or "UNMAPPED",
                        internal_sku=row.get("SKU") or None,
                        product_name=row.get("商品名称") or "未命名商品",
                        specification=row.get("规格") or "",
                        quantity=int(row.get("数量") or 1),
                        unit_price=amount,
                        barcode=row.get("条码") or "",
                    ),
                ),
            )
            _, was_created = upsert_unified(db, principal.tenant_id, incoming)
            created += int(was_created)
        except (KeyError, ValueError, InvalidOperation) as exc:
            errors.append({"line": line, "message": str(exc)})
    db.commit()
    return {"rows": len(rows), "created": created, "errors": errors[:100]}
